from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from utils import poner_borde_inferior

# ==============================================================================
# GESTOR DE WORD - v0.9.7.4m
# ==============================================================================

class GestorWord:
    """ Clase para encapsular la lógica de creación de documentos Word (.docx). """
    
    @staticmethod
    def crear_documento(ruta, datos_json, autor_nombre, historial_coloquio=None):
        """ Crea y guarda el archivo Word en la ruta especificada. """
        doc = Document()
        
        doc.add_paragraph("\n\n\n") 
        
        titulo = doc.add_paragraph()
        run_t = titulo.add_run("ANÁLISIS SEMIÓTICO")
        run_t.font.size = Pt(26)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(0, 51, 102)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph("\n")
        
        subtitulo = doc.add_paragraph()
        run_s = subtitulo.add_run(f"Crítico Invitado: {autor_nombre}")
        run_s.font.size = Pt(14)
        run_s.italic = True
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        fecha = doc.add_paragraph()
        fecha_str = datetime.now().strftime("%d de %B de %Y")
        run_f = fecha.add_run(f"Fecha: {fecha_str}")
        run_f.font.size = Pt(10)
        run_f.font.color.rgb = RGBColor(120, 120, 120)
        fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_page_break()

        def agregar_seccion(titulo_texto, contenido_texto):
            p_tit = doc.add_paragraph()
            run_tit = p_tit.add_run(titulo_texto.upper())
            run_tit.font.size = Pt(14)
            run_tit.font.bold = True
            run_tit.font.color.rgb = RGBColor(50, 50, 50)
            poner_borde_inferior(p_tit)
            doc.add_paragraph() 
            p_cont = doc.add_paragraph(contenido_texto)
            p_cont.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            doc.add_paragraph("\n")

        if 'analisis_autor' in datos_json:
            agregar_seccion("La Mirada Crítica", datos_json['analisis_autor'])
        
        if 'sinopsis' in datos_json:
            agregar_seccion("Sinopsis Narrativa", datos_json['sinopsis'])

        if 'titulos' in datos_json:
            p_t = doc.add_paragraph()
            run_t = p_t.add_run("TÍTULOS ALTERNATIVOS")
            run_t.font.size = Pt(14)
            run_t.font.bold = True
            poner_borde_inferior(p_t)
            
            for t in datos_json['titulos']:
                doc.add_paragraph(f"• {t}", style='List Bullet')
            doc.add_paragraph("\n")

        if 'conceptos' in datos_json:
            p_c = doc.add_paragraph()
            run_c = p_c.add_run("CONCEPTOS CLAVE")
            run_c.font.size = Pt(14)
            run_c.font.bold = True
            poner_borde_inferior(p_c)
            
            conceptos_str = " | ".join([c.upper() for c in datos_json['conceptos']])
            p_cloud = doc.add_paragraph(conceptos_str)
            p_cloud.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cloud.runs[0].font.bold = True
            p_cloud.runs[0].font.color.rgb = RGBColor(100, 100, 100)

        if historial_coloquio:
            doc.add_page_break()
            p_col = doc.add_paragraph()
            run_col = p_col.add_run("COLOQUIO FINAL CON EL AUTOR")
            run_col.font.size = Pt(14)
            run_col.font.bold = True
            run_col.font.color.rgb = RGBColor(150, 50, 50)
            poner_borde_inferior(p_col)
            doc.add_paragraph("\n")

            for remitente, mensaje in historial_coloquio:
                p_msg = doc.add_paragraph()
                nombre_mostrar = autor_nombre.upper() if remitente == "Autor" else "USUARIO"
                
                run_nom = p_msg.add_run(f"{nombre_mostrar}: ")
                run_nom.font.bold = True
                if remitente == "Autor":
                    run_nom.font.color.rgb = RGBColor(0, 0, 100)
                else:
                    run_nom.font.color.rgb = RGBColor(80, 80, 80)
                
                p_msg.add_run(mensaje)
                p_msg.paragraph_format.space_after = Pt(6)
                p_msg.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        doc.save(ruta)
        print(f"[Terminal]: Documento Word generado en -> {ruta}")
